"""Review exports keep every scene, literal text, private images and existing choices."""
import csv
import io
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as appmod
import material_review as review
import material_store as store


@pytest.fixture
def sample(tmp_path, monkeypatch):
    monkeypatch.setattr(appmod, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(appmod, "APP_PASSWORD", "test-password")
    monkeypatch.delenv("GOOGLE_SHEETS_CLIENT_ID", raising=False)
    d = tmp_path / "sample"
    (d / "images").mkdir(parents=True)
    Image.new("RGB", (640, 360), "red").save(d / "images" / "1.png")
    Image.new("RGB", (400, 600), "blue").save(d / "images" / "2.png")
    rows = [
        {"no":1, "chapter_index":1, "sentence":"価格と供給を比べます。", "route":"diagram", "display":"image", "status":"ok", "filename":"1.png"},
        {"no":2, "chapter_index":1, "sentence":"前の図を続けます。", "route":"diagram", "display":"hold", "status":"hold"},
        {"no":3, "chapter_index":1, "sentence":"=HYPERLINK(\"https://example.invalid\",\"入力原文\")", "route":"skip", "display":"none", "status":"skipped"},
        {"no":4, "chapter_index":2, "sentence":"別の章です。", "route":"diagram", "display":"hold", "status":"hold"},
        {"no":5, "chapter_index":2, "sentence":"生成に失敗しました。", "route":"illustration", "display":"image", "status":"failed", "error":"処理失敗"},
        {"no":6, "chapter_index":2, "sentence":"縦長の候補です。", "route":"illustration", "display":"image", "status":"ok", "filename":"2.png"},
    ]
    store.save(d / "rows_progress.json", {"rows":rows})
    store.save(d / "job.json", {"status":"completed", "channel_id":"keizai"})
    store.save(d / "manifest.json", {"title":"書き出し確認", "rows":[]})
    library = store.load_library(d)
    first = store.snapshot(d, library, 1, d / "images" / "1.png", "採用候補")
    library["selections"]["1"] = first["id"]
    library["flags"]["1"] = "数字を確認"
    store.commit_library(d, library, "fixture")
    client = appmod.app.test_client()
    with client.session_transaction() as session:
        session["authenticated"] = True
        session["material_csrf"] = "test-csrf"
    return d, client


def test_all_rows_adoption_hold_and_skips_are_preserved(sample):
    d, _ = sample
    before = {p:p.read_bytes() for p in d.rglob("*") if p.is_file()}
    title, rows = review.review_rows(d)
    assert title == "書き出し確認"
    assert [r["values"][0] for r in rows] == [1,2,3,4,5,6]
    assert rows[0]["values"][4] == "採用済み"
    assert rows[0]["values"][11] == "数字を確認"
    assert rows[1]["path"] == rows[0]["path"]
    assert rows[2]["path"] is None and rows[2]["values"][4] == "スキップ"
    assert rows[3]["path"] is None and rows[3]["values"][4] == "継続元なし・要確認"
    assert rows[4]["values"][4] == "生成失敗"
    assert rows[5]["values"][4] == "生成済み・未採用"
    assert before == {p:p.read_bytes() for p in d.rglob("*") if p.is_file()}


def test_rendered_progress_scripts_parse_before_running_controls(sample):
    node = shutil.which("node")
    if not node:
        pytest.skip("Node is required for browser script syntax validation")
    _, client = sample
    page = client.get("/progress/sample").get_data(as_text=True)
    scripts = re.findall(r"<script(?:\s[^>]*)?>([\s\S]*?)</script>", page)
    assert scripts
    for script in scripts:
        if script.strip():
            result = subprocess.run([node, "--check", "-"], input=script, capture_output=True, text=True, encoding="utf-8")
            assert result.returncode == 0, result.stderr


def test_xlsx_embeds_images_and_literal_formula_like_manuscript(sample):
    d, client = sample
    response = client.post("/api/material-jobs/sample/review.xlsx", headers={"X-Material-CSRF":"test-csrf"})
    assert response.status_code == 200
    assert response.headers["X-Review-Rows"] == "6"
    assert response.headers["Cache-Control"] == "no-store"
    content = response.data
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert len([n for n in archive.namelist() if n.startswith("xl/media/")]) == 3
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")
        assert b"TargetMode=\"External\"" not in archive.read("xl/drawings/_rels/drawing1.xml.rels")
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    assert sheet.max_row == 7 and sheet.max_column == 12
    assert sheet["C1"].value == "コメント" and sheet["C2"].value in (None, "")
    assert sheet["D4"].value.startswith("=HYPERLINK") and sheet["D4"].data_type == "s"
    assert sheet.freeze_panes == "D2" and sheet.row_dimensions[2].height == 126
    workbook.close()


def test_bad_images_are_visible_and_never_replaced_with_a_wrong_candidate(sample):
    d, _ = sample
    (d / "images" / "2.png").write_bytes(b"broken-image")
    stream = io.BytesIO()
    review.write_workbook(d, stream)
    workbook = load_workbook(stream)
    assert "読み込めません" in workbook.active["B7"].value
    assert "画像要確認" in workbook.active["E7"].value
    workbook.close()


def test_csv_preserves_old_columns_and_appends_current_status(sample):
    d, client = sample
    raw = '章,センテンス,№,画像\n1,元の原稿,1,old.png\n1,スキップ,3,\n'
    (d / "result.csv").write_text(raw, encoding="utf-8-sig")
    response = client.get("/download/csv/sample")
    rows = list(csv.reader(io.StringIO(response.data.decode("utf-8-sig"))))
    assert rows[0][:4] == ["章", "センテンス", "№", "画像"]
    assert rows[1][:4] == ["1", "元の原稿", "1", "old.png"]
    assert rows[0][-4:] == ["画像URL", "生成・採用状態", "要修正メモ", "コメント"]
    assert rows[1][-3:] == ["採用済み", "数字を確認", ""]
    assert rows[2][-3] == "スキップ"
    assert (d / "result.csv").read_text(encoding="utf-8-sig") == raw
    response = client.get("/api/material-jobs/sample/review.csv")
    rows = list(csv.reader(io.StringIO(response.data.decode("utf-8-sig"))))
    assert len(rows) == 7 and rows[3][3].startswith("'=HYPERLINK")


def test_auth_csrf_and_no_unconfigured_google_button(sample):
    _, client = sample
    assert client.post("/api/material-jobs/sample/review.xlsx").status_code == 403
    page = client.get("/progress/sample").text
    assert 'data-sheet-create disabled' in page
    assert '連携設定の準備中' in page
    with client.session_transaction() as session:
        session.clear()
    assert client.get("/api/material-jobs/sample/review.csv").status_code == 401
    assert client.post("/api/material-jobs/sample/review.xlsx").status_code == 401


def test_skipped_scene_gets_a_new_prompt_without_changing_other_scenes(sample, monkeypatch):
    import generator
    import prompter
    import utils
    d, client = sample
    original = store.read(d / "rows_progress.json")["rows"]
    adopted_before = (d / "material_library.json").read_bytes()
    monkeypatch.setattr(appmod, "resolve_channel_keys", lambda _: {"openai":"test", "gemini":"test", "anthropic":"test"})
    monkeypatch.setattr(utils, "get_anthropic_client", lambda *_: object())
    monkeypatch.setattr(prompter, "generate_all_prompts", lambda client, rows, **kw: [{**rows[0], "prompt":"Test educational diagram"}])
    calls = []
    def generate(**kw):
        calls.append(kw["prompts"][0])
        Image.new("RGB", (64, 36), "green").save(kw["output_dir"] / "3.png")
        return [{"success":True, "filename":"3.png"}]
    monkeypatch.setattr(generator, "run_parallel_generation", generate)
    response = client.post("/api/regenerate/sample/3", data={"force_route":"diagram", "extra_instruction":"矢印でつなぐ"})
    assert response.status_code == 200 and response.json["filename"] == "3.png"
    assert len(calls) == 1 and "矢印でつなぐ" in calls[0]["prompt"]
    updated = store.read(d / "rows_progress.json")["rows"]
    assert updated[2]["status"] == "ok" and updated[2]["route"] == "diagram"
    assert original[:2] + original[3:] == updated[:2] + updated[3:]
    assert (d / "material_library.json").read_bytes() == adopted_before


def test_zip_csv_has_comments_without_modifying_saved_csv(sample):
    d, client = sample
    original = b'number,title\r\n1,normal\r\n'
    (d / "result.csv").write_bytes(original)
    response = client.get('/download/sample')
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        rows = list(csv.reader(io.StringIO(archive.read('result.csv').decode('utf-8-sig'))))
        assert rows[0][-1] == 'コメント'
        assert archive.read('images/1.png') == (d / 'images' / '1.png').read_bytes()
    assert (d / "result.csv").read_bytes() == original


def test_google_upload_and_skip_ui_contract():
    node = shutil.which("node")
    if not node:
        pytest.skip("Node is required for UI contract tests")
    subprocess.run([node, "--test", "tests/test_material_sheet_export.cjs"],
                   cwd=Path(__file__).resolve().parents[1], check=True)
